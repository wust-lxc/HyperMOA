"""Shared inputs and metric primitives for the current benchmark evaluator."""

from __future__ import annotations

import sys as _sys
from pathlib import Path as _Path
_sys.path.insert(0, str(_Path(__file__).resolve().parents[3]))


from dataclasses import dataclass
from pathlib import Path

import anndata as ad
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from scipy.spatial import cKDTree, distance


def project_root() -> Path:
    path = Path(__file__).resolve().parent
    for candidate in (path, *path.parents):
        if (candidate / "data").is_dir() and (candidate / "benchmark .xlsx").is_file():
            return candidate
    raise FileNotFoundError("Cannot locate benchmark project root")


ROOT = project_root()
METHODS = ("HyperMOA", "3d-OT", "SLAT", "PASTE", "SANTO")
METRICS = {"PAA", "SMR", "FOSCTTM", "Chamfer Distance", "CI", "PCC", "Macro-F1", "Balanced Accuracy"}
MISSING_LABELS = {"", "na", "n/a", "nan", "none", "null", "unknown"}


@dataclass(frozen=True)
class DatasetSpec:
    name: str
    slices: tuple[str, ...]
    raw_paths: tuple[Path, ...]
    result_root: Path
    workbook_sheet: str
    has_true_pairs: bool
    has_ci: bool


SPECS = {
    "Human_lymph_node": DatasetSpec(
        "Human_lymph_node", ("slice1", "slice2", "slice3"),
        tuple(ROOT / f"data/datasets/Human_lymph_node/slice{i}/s{i}_adata_rna.h5ad" for i in range(1, 4)),
        ROOT / "data/results/alignment/Human_lymph_node",
        "Human_lymph_node(global)", False, False,
    ),
    "Simulation": DatasetSpec(
        "Simulation", tuple(f"Simulation{i}" for i in range(1, 6)),
        tuple(ROOT / f"data/datasets/Simulation/slice{i}/adata_RNA.h5ad" for i in range(1, 6)),
        ROOT / "data/results/alignment/Simulation", "Simulation(global)", True, True,
    ),
    "Mouse_Brain": DatasetSpec(
        "Mouse_Brain", ("Mouse_Brain_ATAC", "Mouse_Brain_H3K27ac"),
        (ROOT / "data/datasets/Mouse_Brain/ATAC/adata_RNA.h5ad",
         ROOT / "data/datasets/Mouse_Brain/H3K27ac/adata_RNA.h5ad"),
        ROOT / "data/results/alignment/Mouse_Brain", "Mouse_Brain", True, False,
    ),
}

HYPERMOA_SOURCE_FILES = {
    "Human_lymph_node": ROOT / "HyperMOA/Human_lymph_node_result/adata.h5ad",
    "Simulation": ROOT / "HyperMOA/Simulation_result/adata.h5ad",
    "Mouse_Brain": ROOT / "HyperMOA/Mouse_brain_result/H3K27ac_ATAC/adata.h5ad",
}


def labels(data: ad.AnnData) -> np.ndarray:
    for key in ("truth", "ground_truth", "final_annot", "lab"):
        if key in data.obs:
            return data.obs[key].astype("string").to_numpy(dtype=object)
    return np.full(data.n_obs, pd.NA, dtype=object)


def valid_labels(values: np.ndarray) -> np.ndarray:
    return np.asarray([
        not pd.isna(value) and str(value).strip().lower() not in MISSING_LABELS
        for value in values
    ], dtype=bool)


def nearest_match(source: np.ndarray, target: np.ndarray) -> np.ndarray:
    return np.asarray(cKDTree(target).query(source, k=1)[1], dtype=int)


def median_nn_radius(coordinates: np.ndarray) -> float:
    distances, _ = cKDTree(coordinates).query(coordinates, k=2)
    return float(np.median(distances[:, 1]))


def chamfer_distance(source: np.ndarray, target: np.ndarray) -> float:
    forward = cKDTree(target).query(source, k=1)[0]
    reverse = cKDTree(source).query(target, k=1)[0]
    return float((forward.mean() + reverse.mean()) / 2.0)


def foscttm_chunked(source: np.ndarray, target: np.ndarray, true_target: np.ndarray,
                    chunk_size: int = 512) -> float:
    fractions = []
    denominator = max(1, len(target) - 1)
    for start in range(0, len(source), chunk_size):
        stop = min(start + chunk_size, len(source))
        pair_distances = distance.cdist(source[start:stop], target)
        truth_distance = pair_distances[np.arange(stop - start), true_target[start:stop]]
        fractions.append((pair_distances < truth_distance[:, None]).sum(axis=1) / denominator)
    return float(np.concatenate(fractions).mean())


def hypermoa_batch_name(dataset: str, slice_name: str) -> str:
    return slice_name.replace("Mouse_Brain_", "") if dataset == "Mouse_Brain" else slice_name


def workbook_alignment_cells(path: Path) -> dict:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sources = {}
    try:
        for spec in SPECS.values():
            sheet = workbook[spec.workbook_sheet]
            method_columns = [column for column in range(1, sheet.max_column + 1)
                              if str(sheet.cell(2, column).value).strip() == "HyperMOA"]
            method_column = max(method_columns)
            metrics = []
            for column in range(method_column + 1, sheet.max_column + 1):
                name = str(sheet.cell(1, column).value or "").replace("↑", "").replace("↓", "").replace("⬇", "").strip()
                name = " ".join(name.split())
                name = {
                    "Chamfer / NN spacing": "Chamfer Distance",
                    "PAA (grid proxy)": "PAA",
                    "FOSCTTM (grid proxy)": "FOSCTTM",
                }.get(name, name)
                if name in METRICS:
                    metrics.append((column, name))
            cells = {}
            for row in range(2, sheet.max_row + 1):
                method = str(sheet.cell(row, method_column).value or "").strip()
                if method in METHODS:
                    cells[method] = {name: f"{get_column_letter(column)}{row}"
                                     for column, name in metrics}
            sources[spec.name] = {"sheet": sheet.title, "cells": cells}
    finally:
        workbook.close()
    return sources
