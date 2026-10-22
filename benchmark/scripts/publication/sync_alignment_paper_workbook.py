#!/usr/bin/env python3
"""Populate the official workbook without touching clustering tables."""

from __future__ import annotations

import sys as _sys
from pathlib import Path as _Path
_sys.path.insert(0, str(_Path(__file__).resolve().parents[3]))


import sys
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from benchmark.scripts.common import fair_benchmark_common as common
from benchmark.scripts.evaluation.alignment.evaluate_alignment_paper import MAIN_METRICS

ROOT = common.ROOT
MAIN = ROOT / "data/evaluation/fair_benchmark/alignment_paper_main_metrics.csv"
PROXY = ROOT / "data/evaluation/fair_benchmark/alignment_mouse_grid_proxy_metrics.csv"
WORKBOOK = ROOT / "benchmark .xlsx"
HEADERS = {
    "Human_lymph_node": ("Chamfer / NN spacing↓", "PCC↑", "CI↑", "Macro-F1↑", "Balanced Accuracy↑"),
    "Simulation": ("PAA↑", "FOSCTTM↓", "Chamfer / NN spacing↓", "CI↑", "PCC↑", "SMR↑"),
    "Mouse_Brain": ("PAA (grid proxy)↑", "FOSCTTM (grid proxy)↓", "Chamfer / NN spacing↓", "PCC↑"),
}


def main() -> None:
    metrics = pd.read_csv(MAIN).set_index(["dataset", "method"])
    proxies = pd.read_csv(PROXY).set_index("method")
    book = load_workbook(WORKBOOK)
    for dataset, spec in common.SPECS.items():
        sheet = book[spec.workbook_sheet]
        for col in range(12, max(17, sheet.max_column + 1)):
            if dataset == "Mouse_Brain" or col >= 12 + len(MAIN_METRICS[dataset]):
                sheet.cell(1, col).value = None
                for row in range(2, sheet.max_row + 1):
                    if str(sheet.cell(row, 11).value or "").strip() in common.METHODS:
                        sheet.cell(row, col).value = None
        for offset, (metric, header) in enumerate(zip(MAIN_METRICS[dataset], HEADERS[dataset])):
            col = 12 + offset
            if metric == "Balanced Accuracy":
                sheet.cell(1, col)._style = copy(sheet.cell(1, col - 1)._style)
                sheet.column_dimensions["P"].width = 23
            sheet.cell(1, col).value = header
            for row in range(2, sheet.max_row + 1):
                method = str(sheet.cell(row, 11).value or "").strip()
                if method in common.METHODS:
                    cell = sheet.cell(row, col)
                    if metric == "Balanced Accuracy":
                        cell._style = copy(sheet.cell(row, col - 1)._style)
                    cell.value = round(float(metrics.loc[(dataset, method), metric]), 3)
                    cell.number_format = "0.000"
    name = "Mouse_Brain_grid_proxy"
    if name in book:
        del book[name]
    sheet = book.create_sheet(name)
    sheet.append(["Method", "PAA (grid proxy)", "SMR (grid proxy)", "FOSCTTM (grid proxy)", "Shared grid IDs"])
    for method in common.METHODS:
        record = proxies.loc[method]
        sheet.append([method, round(float(record.PAA), 3), round(float(record.SMR), 3),
                      round(float(record.FOSCTTM), 3), int(record.n_grid_pairs)])
    for col, width in {"A": 18, "B": 22, "C": 22, "D": 26, "E": 19}.items():
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.font = copy(book[common.SPECS["Mouse_Brain"].workbook_sheet]["L1"].font)
    sheet["A8"] = "Shared barcode/grid position is not a verified same biological spot."
    book.save(WORKBOOK)
    print(f"Updated {WORKBOOK}")


if __name__ == "__main__":
    main()
