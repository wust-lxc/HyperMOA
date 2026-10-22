#!/usr/bin/env python
"""Publish the fixed five-metric clustering panel without rerunning methods."""

from __future__ import annotations

import sys as _sys
from pathlib import Path as _Path
_sys.path.insert(0, str(_Path(__file__).resolve().parents[3]))


import hashlib
import json
from copy import copy
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
AUDIT_DIR = ROOT / "data/evaluation/fair_benchmark"
WORKBOOK = ROOT / "benchmark _Clustering.xlsx"
FAIR_SLICE = AUDIT_DIR / "clustering_slice_metrics_fair.csv"
FAIR_MACRO = AUDIT_DIR / "clustering_macro_metrics_fair.csv"
ADDITIONAL_SLICE = AUDIT_DIR / "additional_clustering/additional_clustering_slice_metrics.csv"
ADDITIONAL_MACRO = AUDIT_DIR / "additional_clustering/additional_clustering_macro_metrics.csv"
METRICS = ("ARI", "AMI", "Moran's I", "PAS", "CHAOS")
LOWER_IS_BETTER = frozenset(("PAS", "CHAOS"))
METHODS = (
    "HyperMOA", "SpatialGlue", "MISO", "CellCharter", "Harmony", "3d-OT", "SpaMosaic",
)
OUTPUT_STEM = "clustering_paper_metrics"

SHEETS = {
    "Human_lymph_node": {
        "slice1": "Human_lymph_node(slice1)",
        "slice2": "Human_lymph_node(slice2)",
        "slice3": "Human_lymph_node(slice3)",
        "global": "Human_lymph_node(global)",
    },
    "Simulation": {
        **{f"Simulation{i}": f"Simulation(slice{i})" for i in range(1, 6)},
        "global": "Simulation(global)",
    },
    "Mouse_Brain": {
        "Mouse_Brain_ATAC": "Mouse_Brain_ATAC",
        "Mouse_Brain_H3K27ac": "Mouse_Brain_H3K27ac",
        "global": "Mouse_Brain",
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def three_decimals(value: float) -> str:
    return format(Decimal(str(value)).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP), "f")


def rank_score(group: pd.DataFrame) -> pd.Series:
    if set(group["method"]) != set(METHODS) or len(group) != len(METHODS):
        raise ValueError("Each rank group must contain exactly the seven compared methods")
    scores = pd.concat(
        [
            ((group[metric].rank(ascending=metric not in LOWER_IS_BETTER, method="average") - 1)
             / (len(METHODS) - 1)).rename(metric)
            for metric in METRICS
        ],
        axis=1,
    )
    return scores.mean(axis=1)


def combine(fair_path: Path, additional_path: Path, keys: list[str]) -> pd.DataFrame:
    fair = pd.read_csv(fair_path)
    additional = pd.read_csv(additional_path)
    merged = fair.loc[:, [*keys, "ARI", "AMI", "Moran's I"]].merge(
        additional.loc[:, [*keys, "PAS", "CHAOS"]],
        on=keys,
        how="inner",
        validate="one_to_one",
    )
    if len(merged) != len(fair) or len(merged) != len(additional):
        raise ValueError(f"Incomplete metric merge: {fair_path} and {additional_path}")
    if not np.isfinite(merged.loc[:, METRICS].to_numpy(dtype=float)).all():
        raise ValueError("Non-finite metric in paper score input")
    group_keys = [key for key in keys if key != "method"]
    scores = pd.Series(index=merged.index, dtype=float)
    for _, group in merged.groupby(group_keys, sort=False):
        scores.loc[group.index] = rank_score(group).to_numpy()
    merged["Aggregate score"] = scores
    return merged.loc[:, [*keys, *METRICS, "Aggregate score"]]


def sync_workbook(slice_metrics: pd.DataFrame, macro_metrics: pd.DataFrame) -> dict:
    workbook = load_workbook(WORKBOOK)
    slice_lookup = slice_metrics.set_index(["dataset", "slice", "method"])
    macro_lookup = macro_metrics.set_index(["dataset", "method"])
    cell_sources: dict[str, dict[str, dict[str, str]]] = {}
    for dataset, mapping in SHEETS.items():
        for slice_name, sheet_name in mapping.items():
            sheet = workbook[sheet_name]
            if sheet.cell(1, 6).value not in (None, "PAS") or sheet.cell(1, 7).value not in (None, "CHAOS"):
                raise ValueError(f"{sheet_name}: target F/G columns are not empty or expected metrics")
            for column, metric in ((6, "PAS"), (7, "CHAOS")):
                target = sheet.cell(1, column)
                target._style = copy(sheet.cell(1, 5)._style)
                target.value = metric
            cell_sources[sheet_name] = {}
            present = set()
            for row in range(2, sheet.max_row + 1):
                method_value = sheet.cell(row, 1).value
                if method_value is None or str(method_value).strip() not in METHODS:
                    continue
                method = str(method_value).strip()
                present.add(method)
                values = (
                    macro_lookup.loc[(dataset, method)]
                    if slice_name == "global"
                    else slice_lookup.loc[(dataset, slice_name, method)]
                )
                for existing_column, metric in ((2, "ARI"), (4, "AMI"), (5, "Moran's I")):
                    workbook_value = float(sheet.cell(row, existing_column).value)
                    if abs(workbook_value - float(values[metric])) > 0.000501:
                        raise ValueError(f"{sheet_name}!{sheet.cell(row, existing_column).coordinate} disagrees with fair CSV")
                cell_sources[sheet_name][method] = {}
                for column, metric in ((6, "PAS"), (7, "CHAOS")):
                    target = sheet.cell(row, column)
                    target._style = copy(sheet.cell(row, 5)._style)
                    target.value = float(three_decimals(float(values[metric])))
                    target.number_format = "0.000"
                    cell_sources[sheet_name][method][metric] = target.coordinate
            if present != set(METHODS):
                raise ValueError(f"{sheet_name}: incomplete method rows")
    workbook.save(WORKBOOK)
    return cell_sources


def main() -> None:
    slices = combine(FAIR_SLICE, ADDITIONAL_SLICE, ["dataset", "slice", "method"])
    macro = combine(FAIR_MACRO, ADDITIONAL_MACRO, ["dataset", "method"])
    before_workbook_hash = sha256(WORKBOOK)
    cell_sources = sync_workbook(slices, macro)
    output_files = {}
    for kind, frame in (("slice", slices), ("macro", macro)):
        for suffix, float_format in (("", "%.9f"), ("_3dec", three_decimals)):
            path = AUDIT_DIR / f"{OUTPUT_STEM}_{kind}{suffix}.csv"
            frame.to_csv(path, index=False, float_format=float_format)
            output_files[f"{kind}{suffix}"] = str(path)
    manifest_path = AUDIT_DIR / f"{OUTPUT_STEM}_manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "created_at_utc": datetime.now(timezone.utc).isoformat(),
                "script": str(Path(__file__).resolve()),
                "metrics": list(METRICS),
                "lower_is_better": sorted(LOWER_IS_BETTER),
                "aggregate_definition": "Unweighted mean of all five normalized ranks, separately within each slice or on equal-slice-weight dataset macro means; worst=0, best=1, average-rank ties. Full precision before presentation rounding.",
                "workbook": str(WORKBOOK),
                "workbook_sha256_before": before_workbook_hash,
                "workbook_sha256_after": sha256(WORKBOOK),
                "source_csv_sha256": {str(path): sha256(path) for path in (FAIR_SLICE, FAIR_MACRO, ADDITIONAL_SLICE, ADDITIONAL_MACRO)},
                "workbook_new_metric_cells": cell_sources,
                "output_files": output_files,
            },
            indent=2,
        ) + "\n",
        encoding="utf-8",
    )
    print(f"Updated workbook: {WORKBOOK}")
    print(f"Slice rows: {len(slices)}; macro rows: {len(macro)}")
    print(f"Manifest: {manifest_path}")


if __name__ == "__main__":
    main()
