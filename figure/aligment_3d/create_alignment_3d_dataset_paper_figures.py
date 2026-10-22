#!/usr/bin/env python
"""Create one publication-ready 3D alignment figure per dataset."""

from __future__ import annotations

import hashlib
import json
import os
from datetime import datetime, timezone
from pathlib import Path

for variable in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
    os.environ[variable] = "1"
os.environ.setdefault("MPLBACKEND", "Agg")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import Ellipse, Rectangle
from matplotlib.font_manager import FontProperties
from matplotlib.textpath import TextPath
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

import create_alignment_3d_metrics_figure as shared


ROOT = shared.ROOT
WORKBOOK = ROOT / "benchmark .xlsx"
OUTPUT_DIR = ROOT / "figure/aligment_3d"
METRICS_PATH = OUTPUT_DIR / "alignment_metrics_5_methods.csv"
MANIFEST_PATH = OUTPUT_DIR / "paper_figure_manifest.json"
FAIR_METRICS_PATH = ROOT / "data/evaluation/fair_benchmark/alignment_paper_main_metrics.csv"

METHODS = ("HyperMOA", "3d-OT", "SLAT", "PASTE", "SANTO")
DATASETS = shared.DATASETS
FIGURE_WIDTH_CM = 18.0
FIGURE_HEIGHT_CM = 11.0
FIGURE_WIDTH = FIGURE_WIDTH_CM / 2.54
FIGURE_HEIGHT = FIGURE_HEIGHT_CM / 2.54
FIGURE_DPI = 600
# Physical dimensions measured from the published clustering metric tables.
TABLE_BUBBLE_WIDTH_PT = 34.057
TABLE_BUBBLE_HEIGHT_PT = 6.477
TABLE_METHOD_PITCH_PT = TABLE_BUBBLE_WIDTH_PT/0.62
TABLE_LABEL_GAP_PT = TABLE_METHOD_PITCH_PT - TABLE_BUBBLE_WIDTH_PT

FILE_STEMS = {
    "Human_lymph_node": "human_lymph_node_3d_alignment_paper",
    "Simulation": "simulation_3d_alignment_paper",
    "Mouse_Brain": "mouse_brain_3d_alignment_paper",
}

EXPECTED_METRICS = {
    "Human_lymph_node": ("Chamfer Distance", "PCC", "CI", "Macro-F1", "Balanced Accuracy"),
    "Simulation": ("PAA", "FOSCTTM", "Chamfer Distance", "CI", "PCC", "SMR"),
    "Mouse_Brain": ("Region ARI", "Region AMI", "Chamfer Distance", "PCC"),
}



def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_metrics(workbook_path: Path):
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    frames = {}
    source_cells = {}
    try:
        for dataset, config in DATASETS.items():
            if dataset == 'Mouse_Brain':
                raw = pd.read_csv(FAIR_METRICS_PATH)
                frame = raw.loc[raw.dataset.eq(dataset)].set_index('method').reindex(METHODS)
                frame = frame[['dataset', *EXPECTED_METRICS[dataset]]].copy()
                frame.loc[:, list(EXPECTED_METRICS[dataset])] = frame.loc[:, list(EXPECTED_METRICS[dataset])].round(3)
                frames[dataset] = frame
                source_cells[dataset] = {'source_csv': str(FAIR_METRICS_PATH),
                                         'metrics': list(EXPECTED_METRICS[dataset])}
                continue
            sheet = workbook[config["sheet"]]
            method_columns = [
                column
                for column in range(1, sheet.max_column + 1)
                if str(sheet.cell(2, column).value).strip() == "HyperMOA"
            ]
            if not method_columns:
                raise ValueError(f"Cannot locate the alignment method table in {sheet.title}")
            method_column = max(method_columns)

            metric_columns = []
            for column in range(method_column + 1, sheet.max_column + 1):
                metric = shared.normalize_metric_header(sheet.cell(1, column).value)
                if metric:
                    metric_columns.append((column, metric))
            discovered_metrics = tuple(metric for _, metric in metric_columns)
            if discovered_metrics != EXPECTED_METRICS[dataset]:
                raise ValueError(
                    f"Unexpected metrics in {sheet.title}: {discovered_metrics}; "
                    f"expected {EXPECTED_METRICS[dataset]}"
                )

            records = []
            dataset_cells = {}
            for row in range(2, sheet.max_row + 1):
                method_value = sheet.cell(row, method_column).value
                if method_value is None:
                    continue
                method = str(method_value).strip()
                if method not in METHODS:
                    continue
                record = {"method": method}
                dataset_cells[method] = {}
                for column, metric in metric_columns:
                    value = sheet.cell(row, column).value
                    record[metric] = float(value) if value is not None else np.nan
                    dataset_cells[method][metric] = f"{get_column_letter(column)}{row}"
                records.append(record)

            frame = pd.DataFrame(records).set_index("method").reindex(METHODS)
            if frame.index.tolist() != list(METHODS) or frame.isna().all(axis=1).any():
                raise ValueError(f"Incomplete five-method metric table in {sheet.title}")
            frame.insert(0, "dataset", dataset)
            frames[dataset] = frame
            source_cells[dataset] = {
                "sheet": sheet.title,
                "method_column": get_column_letter(method_column),
                "metric_cells": dataset_cells,
            }
    finally:
        workbook.close()
    return frames, source_cells


def validate_aggregate(dataset: str, aggregate: pd.Series) -> None:
    if aggregate.index.tolist() != list(METHODS) or not np.isfinite(aggregate.to_numpy()).all():
        raise ValueError(f"Invalid aggregate scores for {dataset}")
    if not aggregate.between(0.0, 1.0).all():
        raise ValueError(f"Aggregate scores out of range for {dataset}")


def full_precision_rank_frame(dataset: str, frame: pd.DataFrame) -> pd.DataFrame:
    metrics = EXPECTED_METRICS[dataset]
    raw = pd.read_csv(FAIR_METRICS_PATH).set_index(["dataset", "method"]).loc[dataset].reindex(METHODS)
    values = raw.loc[:, metrics].to_numpy(dtype=float)
    displayed = frame.loc[:, metrics].to_numpy(dtype=float)
    if not np.isfinite(values).all() or np.max(np.abs(values - displayed)) > 0.000501:
        raise ValueError(f"Rounded workbook metrics disagree with fair CSV: {dataset}")
    result = frame.copy()
    result.loc[:, metrics] = values
    return result


def validate_alignment_input(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    obs, obsm = shared.read_metadata(path, ['slice', 'slice_order', 'truth'], ['aligned_spatial'])
    coordinates = obsm['aligned_spatial']
    if coordinates.shape != (len(obs), 2) or not np.isfinite(coordinates).all():
        raise ValueError(f'{path}: invalid aligned_spatial shape or values')


def draw_metric_table(
    axis, frame, metrics, cell_scores, aggregate,
    show_group_header: bool = True, value_fontsize: float = 6.0,
):
    """Show a compact table with fixed physical bubble dimensions and gaps."""
    axis.set_axis_off()
    axis.set_xlim(0.0, 1.0)
    axis.set_ylim(0.0, 1.0)
    labels = {
        'Region ARI':'Regional ARI', 'Region AMI':'Regional AMI',
        'Chamfer Distance':'Chamfer/NN spacing', 'Balanced Accuracy':'Balanced accuracy',
        'Macro-F1':'Macro-F1', 'PAA':'PAA', 'FOSCTTM':'FOSCTTM',
        'CI':'CI', 'PCC':'PCC', 'SMR':'SMR',
    }
    header_y, top_y, bottom_y = 0.945, 0.875, 0.12
    row_step = (top_y-bottom_y)/(len(metrics)+1)
    position = axis.get_position()
    table_width_pt = position.width * axis.figure.get_figwidth() * 72
    table_height_pt = position.height * axis.figure.get_figheight() * 72
    bubble_width = TABLE_BUBBLE_WIDTH_PT/table_width_pt
    bubble_height = TABLE_BUBBLE_HEIGHT_PT/table_height_pt
    font=FontProperties(family='Arial',size=6.6)
    label_strings=[labels.get(metric,metric) for metric in metrics]+['Aggregate score']
    label_width_pt=max(TextPath((0,0),text,prop=font).get_extents().width for text in label_strings)
    padding_pt=4.0
    method_block_pt=(len(METHODS)-1)*TABLE_METHOD_PITCH_PT+TABLE_BUBBLE_WIDTH_PT
    content_width_pt=2*padding_pt+label_width_pt+TABLE_LABEL_GAP_PT+method_block_pt
    left_pt=(table_width_pt-content_width_pt)/2
    label_right=(left_pt+padding_pt+label_width_pt)/table_width_pt
    first_center_pt=label_right*table_width_pt+TABLE_LABEL_GAP_PT+TABLE_BUBBLE_WIDTH_PT/2
    centers={method:(first_center_pt+i*TABLE_METHOD_PITCH_PT)/table_width_pt
             for i,method in enumerate(METHODS)}
    rule_limits=[left_pt/table_width_pt,(left_pt+content_width_pt)/table_width_pt]

    axis.text(label_right,header_y,'Metric',transform=axis.transAxes,
              ha='right',va='center',fontsize=6.8,color='#263238')
    for method in METHODS:
        axis.text(centers[method],header_y,method,transform=axis.transAxes,
                  ha='center',va='center',fontsize=6.8,color='#263238')
    axis.plot(rule_limits,[top_y,top_y],transform=axis.transAxes,
              color='#68767d',lw=0.65)

    def display_value(value, aggregate_value=False):
        return shared.score_value_text(value) if aggregate_value else shared.metric_value_text(value)

    for row_index,metric in enumerate([*metrics,'Aggregate score']):
        y=top_y-(row_index+0.5)*row_step
        is_aggregate = metric=='Aggregate score'
        label=metric if is_aggregate else labels.get(metric,metric)
        if is_aggregate:
            line_y=top_y-row_index*row_step
            axis.plot(rule_limits,[line_y,line_y],transform=axis.transAxes,
                      color='#8a969c',lw=0.55)
        axis.text(label_right,y,label,transform=axis.transAxes,ha='right',va='center',
                  fontsize=6.6,color='#263238')
        for method in METHODS:
            value = float(aggregate.loc[method]) if is_aggregate else frame.loc[method,metric]
            score = float(aggregate.loc[method]) if is_aggregate else float(cell_scores.loc[method,metric])
            x=centers[method]
            if pd.isna(value):
                axis.text(x,y,'-',transform=axis.transAxes,ha='center',va='center',
                          fontsize=value_fontsize,color='#7d8c94')
                continue
            text=display_value(value,is_aggregate)
            color=plt.get_cmap('Blues')(0.22+0.68*score)
            axis.add_patch(Ellipse((x,y),width=bubble_width,height=bubble_height,
                                  transform=axis.transAxes,facecolor=color,edgecolor='none',zorder=2))
            luminance=0.2126*color[0]+0.7152*color[1]+0.0722*color[2]
            axis.text(x,y,text,transform=axis.transAxes,ha='center',va='center',
                      fontsize=value_fontsize,color='white' if luminance<0.58 else '#263238',
                      weight='bold' if score>=0.56 else 'normal',zorder=3)
    axis.plot(rule_limits,[bottom_y,bottom_y],transform=axis.transAxes,
              color='#68767d',lw=0.65)


def build_metric_export(metric_frames, aggregate_scores):
    rows = []
    for dataset, frame in metric_frames.items():
        metrics = [column for column in frame.columns if column != "dataset"]
        for method in METHODS:
            row = {"dataset": dataset, "method": method}
            for metric in shared.METRIC_DIRECTION:
                row[metric] = frame.loc[method, metric] if metric in metrics else np.nan
            row["aggregate_score"] = float(shared.score_value_text(aggregate_scores[dataset].loc[method]))
            rows.append(row)
    return pd.DataFrame(rows)


def render_dataset(
    dataset: str,
    config,
    frame: pd.DataFrame,
    cell_scores: pd.DataFrame,
    aggregate: pd.Series,
):
    stem = FILE_STEMS[dataset]
    png_path = OUTPUT_DIR / f"{stem}.png"
    pdf_path = OUTPUT_DIR / f"{stem}.pdf"
    preview_path = OUTPUT_DIR / f"{stem}_preview.png"

    categories, palette = shared.truth_palette(config["root"], dataset)
    fig = plt.figure(figsize=(FIGURE_WIDTH, FIGURE_HEIGHT), facecolor="white")
    outer = fig.add_gridspec(
        3,
        1,
        height_ratios=(2.45, 0.38, 1.65),
        left=0.025,
        right=0.985,
        bottom=0.035,
        top=0.900,
        hspace=0.01,
    )
    panel_grid = outer[0].subgridspec(1, len(METHODS), wspace=0.02)
    panel_manifest = {}
    panel_title_specs = []
    for column, method in enumerate(METHODS):
        subplot_spec = panel_grid[0, column]
        path = config["root"] / method / "alignment_results.h5ad"
        validate_alignment_input(path)
        axis, panel_info = shared.draw_3d_panel(
            fig,
            subplot_spec,
            path,
            dataset,
            method,
            categories,
            palette,
            show_title=True,
        )
        axis.set_title("")
        panel_title_specs.append((method, subplot_spec))
        panel_manifest[method] = panel_info

    legend_axis = fig.add_subplot(outer[1])
    shared.draw_truth_legend(legend_axis, categories, palette, dataset)
    legend_bounds = legend_axis.get_position()
    legend_axis.set_position(
        [
            legend_bounds.x0,
            legend_bounds.y0 + 0.043,
            legend_bounds.width,
            legend_bounds.height,
        ]
    )

    metrics = [column for column in frame.columns if column != "dataset"]
    table_axis = fig.add_subplot(outer[2])
    draw_metric_table(table_axis, frame, metrics, cell_scores, aggregate)

    for method, subplot_spec in panel_title_specs:
        bounds = subplot_spec.get_position(fig)
        fig.text(
            (bounds.x0 + bounds.x1) / 2,
            bounds.y1 + 0.006,
            method,
            ha="center",
            va="bottom",
            fontsize=10.5,
            weight="bold",
            zorder=100,
        )

    fig.savefig(png_path, dpi=FIGURE_DPI, facecolor="white")
    fig.savefig(pdf_path, dpi=FIGURE_DPI, facecolor="white")
    plt.close(fig)

    with Image.open(png_path) as image:
        pixel_size = list(image.size)
        dpi = image.info.get("dpi")
        preview = image.convert("RGB")
        preview.thumbnail((1800, 720), Image.Resampling.LANCZOS)
        preview.save(preview_path, optimize=True)
    expected_pixels = [int(FIGURE_WIDTH * FIGURE_DPI), int(FIGURE_HEIGHT * FIGURE_DPI)]
    if pixel_size != expected_pixels:
        raise ValueError(f"Unexpected PNG dimensions for {dataset}: {pixel_size}")

    return {
        "size_inches": {"width": FIGURE_WIDTH, "height": FIGURE_HEIGHT},
        "dpi": dpi,
        "pixel_size": pixel_size,
        "panels": panel_manifest,
        "outputs": {
            "png": str(png_path),
            "pdf": str(pdf_path),
            "preview": str(preview_path),
        },
    }


def main():
    if not WORKBOOK.exists():
        raise FileNotFoundError(WORKBOOK)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    metric_frames, source_cells = read_metrics(WORKBOOK)
    metric_details = {}
    aggregate_scores = {}
    for dataset, frame in metric_frames.items():
        metrics, cell_scores, aggregate = shared.rank_scores(full_precision_rank_frame(dataset, frame))
        if tuple(metrics) != EXPECTED_METRICS[dataset]:
            raise ValueError(f"Metric order changed for {dataset}: {metrics}")
        validate_aggregate(dataset, aggregate)
        metric_details[dataset] = cell_scores
        aggregate_scores[dataset] = aggregate

    metric_export = build_metric_export(metric_frames, aggregate_scores)
    metric_export.to_csv(METRICS_PATH, index=False, float_format="%.3f")
    if len(metric_export) != len(DATASETS) * len(METHODS):
        raise ValueError("Unexpected number of metric export rows")

    figures = {}
    for dataset, config in DATASETS.items():
        figures[dataset] = render_dataset(
            dataset,
            config,
            metric_frames[dataset],
            metric_details[dataset],
            aggregate_scores[dataset],
        )

    manifest = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "workbook": {
            "path": str(WORKBOOK),
            "sha256": sha256(WORKBOOK),
            "source_cells": source_cells,
        },
        "method_order": list(METHODS),
        "dataset_order": list(DATASETS),
        "figure_design": {
            "size_cm": {"width": FIGURE_WIDTH_CM, "height": FIGURE_HEIGHT_CM},
            "size_inches": {"width": FIGURE_WIDTH, "height": FIGURE_HEIGHT},
            "dpi": FIGURE_DPI,
            "layout": "five 3D panels, ground-truth legend, metric table below",
            "xy_plot_normalization": "independent x/y min-max normalization within each slice",
            "z_coordinate": "negative slice_order with spacing 0.48; not physical thickness",
            "color": "ground-truth annotation, consistent within each dataset",
            "aggregate_score": (
                "mean of 1-(average_rank-1)/(valid_method_count-1), "
                "recomputed across the five displayed methods"
            ),
        },
        "rank_input": str(FAIR_METRICS_PATH) + " (full precision; workbook values displayed to 3 decimals)",
        "mouse_main_metrics": list(EXPECTED_METRICS["Mouse_Brain"]),
        "mouse_aggregate_score": "equal-weight mean of normalized ranks of regional ARI, regional AMI, Chamfer and PCC",
        "mouse_region_protocol": str(ROOT / "data/evaluation/fair_benchmark/alignment_mouse_region_manifest.json"),
        "mouse_region_caveat": "Grid proxies are supplementary; mouse main metrics use supplied regional partitions.",
        "mouse_two_vs_four_metric_sensitivity": str(ROOT / "data/evaluation/fair_benchmark/alignment_mouse_score_sensitivity.csv"),
        "metrics_csv": str(METRICS_PATH),
        "figures": figures,
    }
    serialized = json.dumps(manifest, indent=2, ensure_ascii=True) + "\n"
    MANIFEST_PATH.write_text(serialized, encoding="utf-8")

    print(f"Metrics: {METRICS_PATH}")
    print(f"Manifest: {MANIFEST_PATH}")
    for dataset, details in figures.items():
        outputs = details["outputs"]
        print(
            f"{dataset}: {outputs['png']} "
            f"({details['pixel_size'][0]}x{details['pixel_size'][1]})"
        )


if __name__ == "__main__":
    main()
