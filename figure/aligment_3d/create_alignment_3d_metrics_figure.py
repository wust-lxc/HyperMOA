#!/usr/bin/env python
"""Create the final 18 x 22 inch 3D alignment and metric comparison figure."""

from __future__ import annotations

import hashlib
import json
import math
import os
import sys
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

for variable in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
    os.environ[variable] = "1"
os.environ.setdefault("MPLBACKEND", "Agg")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from matplotlib.patches import Rectangle
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from benchmark.scripts.common.h5_annotations import read_metadata
WORKBOOK = ROOT / "benchmark .xlsx"
OUTPUT_DIR = ROOT / "figure/aligment_3d/overview"
PNG_PATH = OUTPUT_DIR / "alignment_3d_metrics_composite_18x22.png"
PDF_PATH = OUTPUT_DIR / "alignment_3d_metrics_composite_18x22.pdf"
PREVIEW_PATH = OUTPUT_DIR / "alignment_3d_metrics_composite_preview.png"
METRIC_CSV_PATH = OUTPUT_DIR / "final_alignment_metrics.csv"
MANIFEST_PATH = OUTPUT_DIR / "figure_source_manifest.json"

METHODS = ("HyperMOA", "3d-OT", "SLAT", "PASTE", "SANTO")
DATASETS = {
    "Human_lymph_node": {
        "display": "Human lymph node",
        "sheet": "Human_lymph_node(global)",
        "root": ROOT / "data/results/alignment/Human_lymph_node",
    },
    "Simulation": {
        "display": "Simulation",
        "sheet": "Simulation(global)",
        "root": ROOT / "data/results/alignment/Simulation",
    },
    "Mouse_Brain": {
        "display": "Mouse Brain",
        "sheet": "Mouse_Brain",
        "root": ROOT / "data/results/alignment/Mouse_Brain",
    },
}

METRIC_DIRECTION = {
    "Region ARI": "higher",
    "Region AMI": "higher",
    "PAA": "higher",
    "FOSCTTM": "lower",
    "Chamfer Distance": "lower",
    "CI": "higher",
    "Macro-F1": "higher",
    "Balanced Accuracy": "higher",
    "PCC": "higher",
    "SMR": "higher",
}
METRIC_LABEL = {
    "Region ARI": "Region\nARI\u2191",
    "Region AMI": "Region\nAMI\u2191",
    "PAA": "PAA\u2191",
    "FOSCTTM": "FOSCTTM\u2193",
    "Chamfer Distance": "Chamfer / NN\nspacing\u2193",
    "CI": "CI\u2191",
    "Macro-F1": "Macro-F1\u2191",
    "Balanced Accuracy": "Balanced\naccuracy\u2191",
    "PCC": "PCC\u2191",
    "SMR": "SMR\u2191",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_metric_header(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    for marker in ("\u2b07", "\u2193", "\u2191"):
        text = text.replace(marker, "")
    text = " ".join(text.split())
    aliases = {
        "Region ARI": "Region ARI",
        "Region AMI": "Region AMI",
        "PAA": "PAA",
        "PAA (grid proxy)": "PAA",
        "FOSCTTM": "FOSCTTM",
        "FOSCTTM (grid proxy)": "FOSCTTM",
        "Chamfer Distance": "Chamfer Distance",
        "Chamfer / NN spacing": "Chamfer Distance",
        "CI": "CI",
        "Macro-F1": "Macro-F1",
        "Balanced Accuracy": "Balanced Accuracy",
        "PCC": "PCC",
        "SMR": "SMR",
    }
    return aliases.get(text)


def read_alignment_metrics(workbook_path: Path):
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    frames = {}
    source_cells = {}
    for dataset, config in DATASETS.items():
        sheet = workbook[config["sheet"]]
        method_columns = [
            column
            for column in range(1, sheet.max_column + 1)
            if str(sheet.cell(2, column).value).strip() == "HyperMOA"
        ]
        method_column = max(method_columns)
        metric_columns = []
        for column in range(method_column + 1, sheet.max_column + 1):
            metric = normalize_metric_header(sheet.cell(1, column).value)
            if metric:
                metric_columns.append((column, metric))
        if not metric_columns:
            raise ValueError(f"No alignment metrics found in {sheet.title}")

        rows = []
        cell_map = {}
        for row in range(2, sheet.max_row + 1):
            method_value = sheet.cell(row, method_column).value
            if method_value is None:
                continue
            method = str(method_value).strip()
            if method not in METHODS:
                continue
            record = {"method": method}
            cell_map[method] = {}
            for column, metric in metric_columns:
                value = sheet.cell(row, column).value
                record[metric] = float(value) if value is not None else np.nan
                cell_map[method][metric] = f"{get_column_letter(column)}{row}"
            rows.append(record)
        frame = pd.DataFrame(rows).set_index("method").reindex(METHODS)
        frame.insert(0, "dataset", dataset)
        frames[dataset] = frame
        source_cells[dataset] = {
            "sheet": sheet.title,
            "method_column": get_column_letter(method_column),
            "metric_cells": cell_map,
        }
    workbook.close()
    return frames, source_cells


def rank_scores(frame: pd.DataFrame):
    metrics = [column for column in frame.columns if column != "dataset"]
    cell_scores = pd.DataFrame(index=frame.index, columns=metrics, dtype=float)
    for metric in metrics:
        values = frame[metric].astype(float)
        valid = values.notna()
        n_valid = int(valid.sum())
        if n_valid == 0:
            continue
        ascending = METRIC_DIRECTION[metric] == "lower"
        ranks = values[valid].rank(method="average", ascending=ascending)
        if n_valid == 1:
            scores = pd.Series(1.0, index=ranks.index)
        else:
            scores = 1.0 - (ranks - 1.0) / (n_valid - 1.0)
        cell_scores.loc[scores.index, metric] = scores

    aggregate = cell_scores.mean(axis=1, skipna=False)
    return metrics, cell_scores, aggregate


def normalize_truth(values: np.ndarray, dataset: str) -> np.ndarray:
    output = np.asarray(values, dtype=object).copy()
    if dataset == "Human_lymph_node":
        aliases = {
            "medulla cords": "medulla cords",
            "medulla vessel": "medulla vessel",
            "medulla vessels": "medulla vessel",
        }
        for index, value in enumerate(output):
            if pd.isna(value):
                continue
            normalized = str(value).strip().lower()
            output[index] = aliases.get(normalized, str(value).strip())
    return output


def natural_key(value: str):
    import re

    return [int(token) if token.isdigit() else token.lower() for token in re.split(r"(\d+)", str(value))]


def truth_palette(root: Path, dataset: str):
    source_path = root / "HyperMOA/alignment_results.h5ad"
    obs, _ = read_metadata(source_path, ['truth'])
    truth = normalize_truth(obs['truth'].astype('string').to_numpy(), dataset)
    categories = sorted(
        {str(value) for value in truth if pd.notna(value) and str(value) != "nan"},
        key=natural_key,
    )
    if dataset == "Simulation":
        colors = ["#0072B2", "#E69F00", "#009E73", "#D55E00", "#CC79A7"]
    elif len(categories) <= 10:
        colors = list(plt.get_cmap("tab10").colors)
    else:
        base_colors = list(plt.get_cmap("tab20").colors) + list(plt.get_cmap("Set3").colors)
        colors = [tuple(np.clip(np.asarray(color) * 0.78, 0.0, 1.0)) for color in base_colors]
    palette = {category: colors[index % len(colors)] for index, category in enumerate(categories)}
    palette["NA"] = "#bdbdbd"
    return categories, palette


def minmax_per_slice(coordinates: np.ndarray, slices: np.ndarray):
    output = np.zeros_like(coordinates, dtype=float)
    for slice_name in pd.unique(slices):
        mask = slices == slice_name
        points = np.asarray(coordinates[mask], dtype=float)
        minimum = points.min(axis=0)
        span = points.max(axis=0) - minimum
        span[span <= 1e-12] = 1.0
        output[mask] = (points - minimum) / span
    return output


def point_size(n_obs: int) -> float:
    if n_obs <= 7000:
        return 4.50
    if n_obs <= 13000:
        return 2.00
    return 1.20


def draw_3d_panel(
    fig,
    subplot_spec,
    path: Path,
    dataset: str,
    method: str,
    categories,
    palette,
    show_title: bool,
    show_coordinate_frame: bool = True,
):
    if not path.exists():
        axis = fig.add_subplot(subplot_spec, projection="3d")
        axis.set_axis_off()
        axis.text2D(
            0.5,
            0.52,
            "Not completed",
            transform=axis.transAxes,
            ha="center",
            va="center",
            fontsize=9,
            color="#66747b",
            weight="semibold",
            bbox={
                "boxstyle": "square,pad=0.8",
                "facecolor": "#f5f7f8",
                "edgecolor": "#c8d0d4",
                "linewidth": 0.7,
            },
        )
        axis.text2D(
            0.5,
            0.42,
            "No final coordinates",
            transform=axis.transAxes,
            ha="center",
            va="center",
            fontsize=7,
            color="#87949a",
        )
        if show_title:
            axis.set_title(method, fontsize=11, weight="bold", pad=11)
        return axis, {"status": "missing", "path": str(path)}

    obs, obsm = read_metadata(path, ['slice', 'slice_order', 'truth'], ['aligned_spatial'])
    coordinates = obsm['aligned_spatial']
    slices = obs['slice'].astype(str).to_numpy()
    orders = obs['slice_order'].astype(int).to_numpy()
    truth = normalize_truth(obs['truth'].astype('string').to_numpy(), dataset)
    n_obs = len(obs)
    n_slices = int(pd.Series(slices).nunique())

    xy = minmax_per_slice(coordinates, slices)
    z = -orders.astype(float) * 0.48
    values = np.asarray(
        ["NA" if pd.isna(value) or str(value) == "nan" else str(value) for value in truth]
    )
    plot_categories = list(categories)
    if np.any(values == "NA"):
        plot_categories.append("NA")

    axis = fig.add_subplot(subplot_spec, projection="3d")
    for category in plot_categories:
        mask = values == category
        if not np.any(mask):
            continue
        axis.scatter(
            xy[mask, 0],
            xy[mask, 1],
            z[mask],
            s=point_size(n_obs),
            c=[palette[category]],
            linewidths=0,
            alpha=1.0,
            depthshade=False,
            rasterized=True,
        )

    if dataset == "Simulation":
        axis.set_xlim(-0.12, 1.12)
        axis.set_ylim(-0.12, 1.12)
        panel_zoom = 1.03
        panel_z_aspect = 0.85
    else:
        axis.set_xlim(0.0, 1.0)
        axis.set_ylim(0.0, 1.0)
        panel_zoom = 1.26
        panel_z_aspect = max(0.55, 0.25 * n_slices)
    z_min = float(z.min())
    axis.set_zlim(z_min - 0.03, 0.03)
    axis.set_xticks([0.0, 0.5, 1.0])
    axis.set_yticks([0.0, 0.5, 1.0])
    axis.set_zticks(sorted(pd.unique(z)))
    axis.set_xticklabels([])
    axis.set_yticklabels([])
    axis.set_zticklabels([])
    axis.tick_params(pad=-2, length=0)
    axis.view_init(elev=27, azim=-58)
    axis.set_box_aspect(
        (1.0, 1.0, panel_z_aspect),
        zoom=panel_zoom,
    )
    if show_coordinate_frame:
        axis.xaxis.pane.set_facecolor((0.97, 0.98, 0.98, 1.0))
        axis.yaxis.pane.set_facecolor((0.97, 0.98, 0.98, 1.0))
        axis.zaxis.pane.set_facecolor((0.97, 0.98, 0.98, 1.0))
        axis.grid(True, alpha=0.22)
    else:
        axis.set_axis_off()
    if show_title:
        axis.set_title(method, fontsize=11, weight="bold", pad=11)
    return axis, {
        "status": "complete",
        "path": str(path),
        "n_obs": n_obs,
        "n_slices": n_slices,
    }


def draw_truth_legend(axis, categories, palette, dataset: str):
    axis.set_axis_off()
    if dataset == "Mouse_Brain":
        for index, subset in enumerate((
            [category for category in categories if not category.startswith("R")],
            [category for category in categories if category.startswith("R")],
        )):
            handles = [Line2D([0], [0], marker="o", linestyle="",
                              markerfacecolor=palette[category], markeredgecolor="none",
                              markersize=4.2, label=category) for category in subset]
            legend = axis.legend(handles=handles, loc="upper center" if index == 0 else "lower center",
                                 bbox_to_anchor=(0.5, 1.0 if index == 0 else 0.0),
                                 ncol=len(subset), frameon=False, fontsize=6.5,
                                 handletextpad=0.25, columnspacing=0.65, borderaxespad=0)
            if index == 0:
                axis.add_artist(legend)
        return
    handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            linestyle="",
            markerfacecolor=palette[category],
            markeredgecolor="none",
            markersize=4.2,
            label=category,
        )
        for category in categories
    ]
    if dataset == "Human_lymph_node":
        n_columns, font_size = 5, 6.8
    else:
        n_columns, font_size = len(categories), 7.2
    axis.legend(
        handles=handles,
        loc="center",
        ncol=n_columns,
        frameon=False,
        fontsize=font_size,
        handletextpad=0.25,
        columnspacing=0.8,
        borderaxespad=0,
    )


def metric_value_text(value: float) -> str:
    return "-" if pd.isna(value) else f"{float(value):.3f}"


def score_value_text(value: float) -> str:
    return str(Decimal(str(value)).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


def draw_metric_table(axis, frame, metrics, cell_scores, aggregate, show_group_header: bool):
    axis.set_axis_off()
    axis.set_xlim(0.0, 1.0)
    axis.set_ylim(0.0, 1.0)

    method_right = 0.215
    aggregate_left = 0.785
    metric_left = method_right + 0.015
    metric_right = aggregate_left - 0.025
    metric_width = (metric_right - metric_left) / max(1, len(metrics))
    metric_centers = {
        metric: metric_left + (index + 0.5) * metric_width
        for index, metric in enumerate(metrics)
    }

    group_y = 0.958
    group_line_y = 0.914
    header_y = 0.852 if show_group_header else 0.930
    top_rule_y = 0.795 if show_group_header else 0.868
    row_height = 0.078
    first_center_y = 0.732 if show_group_header else 0.805
    row_centers = {
        method: first_center_y - index * row_height
        for index, method in enumerate(METHODS)
    }
    bottom_rule_y = first_center_y - (len(METHODS) - 0.5) * row_height

    axis.text(0.015, header_y, "Method", transform=axis.transAxes, ha="left", va="center", fontsize=10, weight="bold")
    if show_group_header:
        axis.text(
            (metric_left + metric_right) / 2,
            group_y,
            "Global\nalignment",
            transform=axis.transAxes,
            ha="center",
            va="center",
            fontsize=10,
            weight="bold",
            linespacing=0.9,
        )
        axis.text(
            (aggregate_left + 0.99) / 2,
            group_y,
            "Aggregate\nscore",
            transform=axis.transAxes,
            ha="center",
            va="center",
            fontsize=10,
            weight="bold",
            linespacing=0.9,
        )
        axis.plot([metric_left, metric_right], [group_line_y, group_line_y], transform=axis.transAxes, color="#bcc6ca", lw=0.8)
        axis.plot([aggregate_left, 0.99], [group_line_y, group_line_y], transform=axis.transAxes, color="#bcc6ca", lw=0.8)
    separator_top = group_line_y if show_group_header else 0.965
    axis.plot([aggregate_left - 0.012, aggregate_left - 0.012], [bottom_rule_y, separator_top], transform=axis.transAxes, color="#8f999e", lw=0.7)

    for metric in metrics:
        axis.text(
            metric_centers[metric],
            header_y,
            METRIC_LABEL[metric],
            transform=axis.transAxes,
            ha="center",
            va="center",
            fontsize=8.2 if metric != "Chamfer Distance" else 7.6,
            color="#354147",
            linespacing=0.9,
        )
    axis.text(
        (aggregate_left + 0.99) / 2,
        header_y,
        "Rank-based\nscore\u2191",
        transform=axis.transAxes,
        ha="center",
        va="center",
        fontsize=8.2,
        color="#354147",
        linespacing=0.9,
    )
    axis.plot([0.0, 0.99], [top_rule_y, top_rule_y], transform=axis.transAxes, color="#4a555a", lw=1.0)

    for index, method in enumerate(METHODS):
        y = row_centers[method]
        axis.text(
            0.015,
            y,
            method,
            transform=axis.transAxes,
            ha="left",
            va="center",
            fontsize=9.2,
            weight="bold",
            color="#263238",
        )

        for metric in metrics:
            value = frame.loc[method, metric]
            score = cell_scores.loc[method, metric]
            x = metric_centers[metric]
            if pd.isna(value):
                axis.text(x, y, "-", transform=axis.transAxes, ha="center", va="center", fontsize=9, color="#8b969b")
                continue
            score_value = float(score)
            bubble_color = plt.get_cmap("Blues")(0.22 + 0.68 * score_value)
            axis.scatter(
                [x],
                [y],
                transform=axis.transAxes,
                s=115 + 355 * score_value,
                c=[bubble_color],
                edgecolors="none",
                zorder=2,
            )
            text_color = "white" if score_value >= 0.56 else "#263238"
            value_text = metric_value_text(value)
            font_size = 7.0 if len(value_text) <= 5 else 6.2
            axis.text(
                x,
                y,
                value_text,
                transform=axis.transAxes,
                ha="center",
                va="center",
                fontsize=font_size,
                weight="bold" if score_value >= 0.56 else "normal",
                color=text_color,
                zorder=3,
            )

        aggregate_value = aggregate.loc[method]
        bar_left = aggregate_left + 0.018
        bar_width = 0.118
        bar_height = 0.035
        if pd.notna(aggregate_value):
            score_value = float(aggregate_value)
            axis.add_patch(
                Rectangle(
                    (bar_left, y - bar_height / 2),
                    bar_width,
                    bar_height,
                    transform=axis.transAxes,
                    facecolor="#e5f0f6",
                    edgecolor="none",
                    zorder=1,
                )
            )
            axis.add_patch(
                Rectangle(
                    (bar_left, y - bar_height / 2),
                    bar_width * score_value,
                    bar_height,
                    transform=axis.transAxes,
                    facecolor="#69acd0",
                    edgecolor="none",
                    zorder=2,
                )
            )
            axis.text(
                0.975,
                y,
                f"{score_value:.3f}",
                transform=axis.transAxes,
                ha="right",
                va="center",
                fontsize=8.1,
                color="#263238",
            )
        else:
            axis.text(
                (aggregate_left + 0.99) / 2,
                y,
                "-",
                transform=axis.transAxes,
                ha="center",
                va="center",
                fontsize=9,
                color="#8b969b",
            )

        if index < len(METHODS) - 1:
            separator_y = y - row_height / 2
            axis.plot(
                [0.0, 0.99],
                [separator_y, separator_y],
                transform=axis.transAxes,
                color="#c7d0d4",
                lw=0.6,
                ls=(0, (1.5, 2.5)),
            )
    axis.plot([0.0, 0.99], [bottom_rule_y, bottom_rule_y], transform=axis.transAxes, color="#4a555a", lw=1.0)


def build_frozen_metric_frame(metric_frames, aggregate_scores):
    rows = []
    for dataset, frame in metric_frames.items():
        metric_columns = [column for column in frame.columns if column != "dataset"]
        for method in METHODS:
            row = {"dataset": dataset, "method": method}
            for metric in METRIC_DIRECTION:
                row[metric] = frame.loc[method, metric] if metric in metric_columns else np.nan
            row["aggregate_score"] = aggregate_scores[dataset].loc[method]
            rows.append(row)
    return pd.DataFrame(rows)


def main():
    if not WORKBOOK.exists():
        raise FileNotFoundError(WORKBOOK)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    metric_frames, source_cells = read_alignment_metrics(WORKBOOK)
    metric_details = {}
    aggregate_scores = {}
    for dataset, frame in metric_frames.items():
        metrics, cell_scores, aggregate = rank_scores(frame)
        metric_details[dataset] = (metrics, cell_scores)
        aggregate_scores[dataset] = aggregate

    frozen = build_frozen_metric_frame(metric_frames, aggregate_scores)
    frozen.to_csv(METRIC_CSV_PATH, index=False, float_format="%.3f")

    fig = plt.figure(figsize=(18, 22), facecolor="white")
    outer = fig.add_gridspec(
        3,
        2,
        width_ratios=(13.45, 4.15),
        left=0.045,
        right=0.993,
        bottom=0.028,
        top=0.958,
        hspace=0.018,
        wspace=0.018,
    )
    fig.suptitle(
        "3D reconstruction and alignment benchmark",
        x=0.52,
        y=0.985,
        fontsize=17,
        weight="bold",
    )

    panel_manifest = {}
    row_specs = []
    for row_index, (dataset, config) in enumerate(DATASETS.items()):
        categories, palette = truth_palette(config["root"], dataset)
        left = outer[row_index, 0].subgridspec(2, 1, height_ratios=(14.0, 0.85), hspace=0.0)
        panel_grid = left[0].subgridspec(1, len(METHODS), wspace=-0.055)
        legend_axis = fig.add_subplot(left[1])
        draw_truth_legend(legend_axis, categories, palette, dataset)
        panel_manifest[dataset] = {}
        for column_index, method in enumerate(METHODS):
            path = config["root"] / method / "alignment_results.h5ad"
            _, panel_info = draw_3d_panel(
                fig,
                panel_grid[0, column_index],
                path,
                dataset,
                method,
                categories,
                palette,
                show_title=row_index == 0,
            )
            panel_manifest[dataset][method] = panel_info

        table_axis = fig.add_subplot(outer[row_index, 1])
        metrics, cell_scores = metric_details[dataset]
        draw_metric_table(
            table_axis,
            metric_frames[dataset],
            metrics,
            cell_scores,
            aggregate_scores[dataset],
            show_group_header=row_index == 0,
        )
        row_specs.append(outer[row_index, 0])

    for dataset, subplot_spec in zip(DATASETS, row_specs):
        bounds = subplot_spec.get_position(fig)
        fig.text(
            0.018,
            (bounds.y0 + bounds.y1) / 2,
            DATASETS[dataset]["display"],
            rotation=90,
            ha="center",
            va="center",
            fontsize=13,
            weight="bold",
            color="#263238",
        )

    for row_index in range(2):
        upper_bounds = outer[row_index, 0].get_position(fig)
        lower_bounds = outer[row_index + 1, 0].get_position(fig)
        separator_y = (upper_bounds.y0 + lower_bounds.y1) / 2
        fig.add_artist(
            Line2D(
                [0.047, 0.987],
                [separator_y, separator_y],
                transform=fig.transFigure,
                color="#d7dee1",
                lw=0.8,
            )
        )

    fig.savefig(PNG_PATH, dpi=300, facecolor="white")
    fig.savefig(PDF_PATH, dpi=300, facecolor="white")
    plt.close(fig)

    with Image.open(PNG_PATH) as image:
        pixel_size = list(image.size)
        preview = image.convert("RGB")
        preview.thumbnail((1620, 1980), Image.Resampling.LANCZOS)
        preview.save(PREVIEW_PATH, optimize=True)
    if pixel_size != [5400, 6600]:
        raise ValueError(f"Unexpected PNG dimensions: {pixel_size}")

    manifest = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "workbook": {
            "path": str(WORKBOOK),
            "sha256": sha256(WORKBOOK),
            "source_cells": source_cells,
        },
        "figure": {
            "size_inches": {"width": 18, "height": 22},
            "dpi": 300,
            "pixel_size": pixel_size,
            "method_order": list(METHODS),
            "dataset_order": list(DATASETS),
            "xy_plot_normalization": "independent x/y min-max normalization within each slice",
            "z_coordinate": "negative slice_order with spacing 0.48; not physical thickness",
            "color": "ground-truth annotation, consistent within each dataset",
            "aggregate_score": "mean of 1-(average_rank-1)/(valid_method_count-1); metric direction from arrows",
        },
        "panels": panel_manifest,
        "outputs": {
            "png": str(PNG_PATH),
            "pdf": str(PDF_PATH),
            "preview": str(PREVIEW_PATH),
            "metrics_csv": str(METRIC_CSV_PATH),
        },
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    print(f"PNG: {PNG_PATH} ({PNG_PATH.stat().st_size} bytes, {pixel_size[0]}x{pixel_size[1]})")
    print(f"PDF: {PDF_PATH} ({PDF_PATH.stat().st_size} bytes)")
    print(f"Preview: {PREVIEW_PATH} ({PREVIEW_PATH.stat().st_size} bytes)")
    print(f"Metrics: {METRIC_CSV_PATH}")
    print(f"Manifest: {MANIFEST_PATH}")


if __name__ == "__main__":
    main()
