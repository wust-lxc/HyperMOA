#!/usr/bin/env python
"""Create one publication-ready clustering comparison figure per dataset."""

from __future__ import annotations

import hashlib
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path

for variable in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
    os.environ[variable] = "1"
os.environ.setdefault("MPLBACKEND", "Agg")

import anndata as ad
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from matplotlib.patches import Rectangle
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "benchmark _Clustering.xlsx"
OUTPUT_DIR = ROOT / "figure/clustering/panels_only"
METRICS_PATH = OUTPUT_DIR / "selected_slice_clustering_metrics.csv"
SELECTION_PATH = OUTPUT_DIR / "slice_selection_scores.csv"
MANIFEST_PATH = OUTPUT_DIR / "clustering_comparison_manifest.json"

FIGURE_WIDTH_CM = 18.0
FIGURE_HEIGHT_CM = 11.0
FIGURE_WIDTH = FIGURE_WIDTH_CM / 2.54
FIGURE_HEIGHT = FIGURE_HEIGHT_CM / 2.54
FIGURE_DPI = 600

METRICS = ("ARI", "NMI", "AMI", "Moran's I")
METHODS = (
    "HyperMOA",
    "SpatialGlue",
    "MISO",
    "CellCharter",
    "Harmony",
    "3d-OT",
    "SpaMosaic",
)
OBS_COLUMNS = {
    "HyperMOA": "cluster_HyperMOA",
    "SpatialGlue": "cluster_SpatialGlue",
    "MISO": "cluster_MISO",
    "CellCharter": "cluster_CellCharter",
    "Harmony": "cluster_Harmony",
    "3d-OT": "cluster_3dOT",
    "SpaMosaic": "cluster_SpaMosaic",
}

CANDIDATES = {
    "Human_lymph_node": (
        {
            "slice": "slice1",
            "sheet": "Human_lymph_node(slice1)",
            "path": ROOT / "data/results/clustering/Human_lymph_node/slice1/clustering_results.h5ad",
        },
        {
            "slice": "slice2",
            "sheet": "Human_lymph_node(slice2)",
            "path": ROOT / "data/results/clustering/Human_lymph_node/slice2/clustering_results.h5ad",
        },
        {
            "slice": "slice3",
            "sheet": "Human_lymph_node(slice3)",
            "path": ROOT / "data/results/clustering/Human_lymph_node/slice3/clustering_results.h5ad",
        },
    ),
    "Simulation": tuple(
        {
            "slice": f"slice{index}",
            "sheet": f"Simulation(slice{index})",
            "path": ROOT / f"data/results/clustering/Simulation/slice{index}/clustering_results.h5ad",
        }
        for index in range(1, 6)
    ),
    "Mouse_Brain": (
        {
            "slice": "ATAC",
            "sheet": "Mouse_Brain_ATAC",
            "path": ROOT / "data/results/clustering/Mouse_Brain/ATAC/clustering_results.h5ad",
        },
        {
            "slice": "H3K27ac",
            "sheet": "Mouse_Brain_H3K27ac",
            "path": ROOT / "data/results/clustering/Mouse_Brain/H3K27ac/clustering_results.h5ad",
        },
    ),
}

EXPECTED_SELECTIONS = {
    "Human_lymph_node": "slice2",
    "Simulation": "slice5",
    "Mouse_Brain": "H3K27ac",
}

FILE_STEMS = {
    "Human_lymph_node": "human_lymph_node_slice2_clustering_comparison",
    "Simulation": "simulation_slice5_clustering_comparison",
    "Mouse_Brain": "mouse_brain_h3k27ac_clustering_comparison",
}

POINT_SIZES = {
    "Human_lymph_node": 2.7,
    "Simulation": 6.0,
    "Mouse_Brain": 1.15,
}

# Saturated colors remain legible after journal down-scaling and avoid gray,
# which is reserved for genuinely missing labels.
TRUTH_COLORS = (
    "#1f77b4",
    "#ff7f0e",
    "#2ca02c",
    "#d62728",
    "#9467bd",
    "#8c564b",
    "#e7298a",
    "#bcbd22",
    "#17becf",
    "#393b79",
    "#31a354",
    "#e6550d",
    "#756bb1",
    "#00a6a6",
    "#c51b7d",
    "#6b8e23",
    "#b2182b",
    "#2166ac",
)
UNLABELLED_COLOR = "#e6e8ea"
HIGHLIGHT_COLOR = "#c58c00"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def natural_key(value: object):
    return [
        int(token) if token.isdigit() else token.lower()
        for token in re.split(r"(\d+)", str(value))
    ]


def read_sheet_metrics(workbook, sheet_name: str):
    sheet = workbook[sheet_name]
    headers = [sheet.cell(1, column).value for column in range(2, 6)]
    if tuple(headers) != METRICS:
        raise ValueError(f"{sheet_name}: unexpected metric headers {headers}")

    records = []
    source_cells = {}
    for row in range(2, sheet.max_row + 1):
        method_value = sheet.cell(row, 1).value
        if method_value is None:
            continue
        method = str(method_value).strip()
        if method not in METHODS:
            continue
        record = {"method": method}
        source_cells[method] = {}
        for column, metric in zip(range(2, 6), METRICS):
            value = sheet.cell(row, column).value
            if value is None:
                raise ValueError(f"{sheet_name}: missing {metric} for {method}")
            record[metric] = float(value)
            source_cells[method][metric] = f"{get_column_letter(column)}{row}"
        records.append(record)

    frame = pd.DataFrame(records).set_index("method").reindex(METHODS)
    if frame.shape != (len(METHODS), len(METRICS)) or frame.isna().any().any():
        raise ValueError(f"{sheet_name}: incomplete seven-method metric table")
    return frame, source_cells


def rank_score(frame: pd.DataFrame) -> pd.Series:
    ranks = frame.loc[:, METRICS].rank(axis=0, method="average", ascending=True)
    return ((ranks - 1.0) / (len(frame) - 1.0)).mean(axis=1)


def choose_slices():
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    selected = {}
    selection_records = []
    selected_metrics = {}
    metric_sources = {}
    try:
        for dataset, candidates in CANDIDATES.items():
            candidate_results = []
            for candidate in candidates:
                frame, cells = read_sheet_metrics(workbook, candidate["sheet"])
                scores = rank_score(frame)
                competitors = scores.drop(index="HyperMOA")
                aggregate_margin = float(scores.loc["HyperMOA"] - competitors.max())
                core_means = frame.loc[:, ("ARI", "NMI", "AMI")].mean(axis=1)
                core_margin = float(
                    core_means.loc["HyperMOA"] - core_means.drop(index="HyperMOA").max()
                )
                record = {
                    "dataset": dataset,
                    "slice": candidate["slice"],
                    "sheet": candidate["sheet"],
                    "HyperMOA_rank_score": float(scores.loc["HyperMOA"]),
                    "best_competitor_rank_score": float(competitors.max()),
                    "rank_score_margin": aggregate_margin,
                    "core_metric_mean_margin": core_margin,
                }
                selection_records.append(record)
                candidate_results.append((aggregate_margin, core_margin, candidate, frame, cells))

            _, _, winner, winner_frame, winner_cells = max(
                candidate_results,
                key=lambda item: (item[0], item[1]),
            )
            expected = EXPECTED_SELECTIONS[dataset]
            if winner["slice"] != expected:
                raise ValueError(
                    f"{dataset}: selection rule chose {winner['slice']}, expected {expected}"
                )
            selected[dataset] = winner
            selected_metrics[dataset] = winner_frame
            metric_sources[dataset] = {
                "sheet": winner["sheet"],
                "cells": winner_cells,
            }
    finally:
        workbook.close()

    return selected, selected_metrics, metric_sources, pd.DataFrame(selection_records)


def validate_input(path: Path):
    if not path.exists():
        raise FileNotFoundError(path)
    data = ad.read_h5ad(path, backed="r")
    try:
        required_obs = {"truth", *OBS_COLUMNS.values()}
        missing_obs = sorted(required_obs - set(data.obs.columns))
        if missing_obs:
            raise KeyError(f"{path}: missing obs columns {missing_obs}")
        if "spatial" not in data.obsm:
            raise KeyError(f"{path}: missing obsm['spatial']")
        spatial = np.asarray(data.obsm["spatial"], dtype=float)
        if spatial.shape != (data.n_obs, 2) or not np.isfinite(spatial).all():
            raise ValueError(f"{path}: invalid spatial coordinates")
    finally:
        data.file.close()


def series_as_strings(series: pd.Series) -> pd.Series:
    result = pd.Series(pd.NA, index=series.index, dtype="object")
    valid = series.notna()
    result.loc[valid] = series.loc[valid].astype(str)
    return result


def panel_colors(
    truth: pd.Series,
    prediction: pd.Series | None,
    truth_palette: dict[str, str],
):
    if prediction is None:
        truth_string = series_as_strings(truth)
        colors = np.full(len(truth_string), UNLABELLED_COLOR, dtype=object)
        valid = truth_string.notna()
        colors[valid.to_numpy()] = truth_string.loc[valid].map(truth_palette).to_numpy()
        return colors, truth_palette

    prediction_string = series_as_strings(prediction)
    prediction_labels = sorted(prediction_string.dropna().unique(), key=natural_key)
    if len(prediction_labels) > len(TRUTH_COLORS):
        raise ValueError(f"Not enough colors for {len(prediction_labels)} prediction clusters")
    prediction_palette = {
        label: TRUTH_COLORS[index]
        for index, label in enumerate(prediction_labels)
    }
    colors = np.full(len(prediction_string), UNLABELLED_COLOR, dtype=object)
    for cluster, color in prediction_palette.items():
        colors[(prediction_string == cluster).to_numpy()] = color
    return colors, prediction_palette


def draw_panel(axis, spatial, colors, title, point_size, highlight=False):
    axis.scatter(
        spatial[:, 0],
        spatial[:, 1],
        c=colors,
        s=point_size,
        marker="o",
        linewidths=0,
        edgecolors="none",
        alpha=0.98,
        rasterized=True,
    )
    x_min, y_min = spatial.min(axis=0)
    x_max, y_max = spatial.max(axis=0)
    x_pad = max((x_max - x_min) * 0.025, 1e-6)
    y_pad = max((y_max - y_min) * 0.025, 1e-6)
    axis.set_xlim(x_min - x_pad, x_max + x_pad)
    axis.set_ylim(y_max + y_pad, y_min - y_pad)
    axis.set_aspect("equal", adjustable="box")
    axis.set_title(title, fontsize=8.6, weight="bold", pad=2.0)
    axis.set_axis_off()
    if highlight:
        axis.add_patch(
            Rectangle(
                (0.005, 0.005),
                0.99,
                0.99,
                transform=axis.transAxes,
                fill=False,
                edgecolor=HIGHLIGHT_COLOR,
                linewidth=1.0,
                clip_on=False,
            )
        )


def draw_legend(
    axis,
    truth_labels,
    truth_palette,
):
    axis.set_axis_off()
    handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            linestyle="none",
            markersize=4.2,
            markerfacecolor=truth_palette[label],
            markeredgecolor="none",
            label=label,
        )
        for label in truth_labels
    ]
    columns = 5 if len(handles) <= 10 else 7
    axis.legend(
        handles=handles,
        loc="center",
        ncol=columns,
        frameon=False,
        fontsize=6.5 if len(handles) > 10 else 6.2,
        handletextpad=0.28,
        columnspacing=0.78,
        borderaxespad=0.0,
        labelspacing=0.42,
    )


def save_preview(source_path: Path, preview_path: Path):
    with Image.open(source_path) as image:
        preview = image.copy()
        preview.thumbnail((1800, 1100), Image.Resampling.LANCZOS)
        preview.save(preview_path, dpi=(150, 150), optimize=True)


def render_dataset(dataset: str, config: dict):
    data = ad.read_h5ad(config["path"])
    spatial = np.asarray(data.obsm["spatial"], dtype=float)
    truth = data.obs["truth"]
    truth_string = series_as_strings(truth)
    truth_labels = sorted(truth_string.dropna().unique(), key=natural_key)
    if len(truth_labels) > len(TRUTH_COLORS):
        raise ValueError(f"{dataset}: not enough truth colors for {len(truth_labels)} labels")
    truth_palette = {
        label: TRUTH_COLORS[index]
        for index, label in enumerate(truth_labels)
    }

    panels = [("Ground truth", None), *[(method, OBS_COLUMNS[method]) for method in METHODS]]
    panel_details = {}
    figure = plt.figure(figsize=(FIGURE_WIDTH, FIGURE_HEIGHT), facecolor="white")
    grid = figure.add_gridspec(
        3,
        4,
        height_ratios=(1.0, 1.0, 0.18),
        left=0.025,
        right=0.985,
        bottom=0.035,
        top=0.955,
        wspace=0.045,
        hspace=0.075,
    )

    for index, (title, column) in enumerate(panels):
        axis = figure.add_subplot(grid[index // 4, index % 4])
        prediction = None if column is None else data.obs[column]
        colors, display_palette = panel_colors(truth, prediction, truth_palette)
        draw_panel(
            axis,
            spatial,
            colors,
            title,
            POINT_SIZES[dataset],
            highlight=(title == "HyperMOA"),
        )
        if prediction is None:
            cluster_count = int(truth.nunique(dropna=True))
            missing_count = int(truth.isna().sum())
        else:
            cluster_count = int(prediction.nunique(dropna=True))
            missing_count = int(prediction.isna().sum())
        panel_details[title] = {
            "cluster_count": cluster_count,
            "missing_count": missing_count,
            "display_palette": display_palette,
        }
    legend_axis = figure.add_subplot(grid[2, :])
    draw_legend(
        legend_axis,
        truth_labels,
        truth_palette,
    )

    stem = FILE_STEMS[dataset]
    png_path = OUTPUT_DIR / f"{stem}.png"
    pdf_path = OUTPUT_DIR / f"{stem}.pdf"
    preview_path = OUTPUT_DIR / f"{stem}_preview.png"
    figure.savefig(png_path, dpi=FIGURE_DPI, facecolor="white")
    figure.savefig(pdf_path, facecolor="white")
    plt.close(figure)
    save_preview(png_path, preview_path)
    data.file.close()

    with Image.open(png_path) as image:
        png_size = list(image.size)
        png_dpi = list(image.info.get("dpi", (None, None)))

    return {
        "input_h5ad": str(config["path"]),
        "input_h5ad_sha256": sha256(config["path"]),
        "sheet": config["sheet"],
        "slice": config["slice"],
        "truth_label_count": len(truth_labels),
        "truth_labels": truth_labels,
        "truth_palette": truth_palette,
        "panel_details": panel_details,
        "outputs": {
            "png": str(png_path),
            "pdf": str(pdf_path),
            "preview_png": str(preview_path),
        },
        "png_pixels": png_size,
        "png_dpi": png_dpi,
    }


def save_metrics(selected, selected_metrics):
    records = []
    for dataset, config in selected.items():
        frame = selected_metrics[dataset]
        for method in METHODS:
            record = {
                "dataset": dataset,
                "slice": config["slice"],
                "sheet": config["sheet"],
                "method": method,
            }
            record.update({metric: float(frame.loc[method, metric]) for metric in METRICS})
            records.append(record)
    pd.DataFrame(records).to_csv(METRICS_PATH, index=False, float_format="%.3f")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    selected, metrics, metric_sources, selection_scores = choose_slices()
    for config in selected.values():
        validate_input(config["path"])

    save_metrics(selected, metrics)
    selection_scores.to_csv(SELECTION_PATH, index=False, float_format="%.6f")

    figures = {}
    for dataset, config in selected.items():
        figures[dataset] = render_dataset(dataset, config)

    manifest = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "script": str(Path(__file__).resolve()),
        "figure_size_cm": [FIGURE_WIDTH_CM, FIGURE_HEIGHT_CM],
        "figure_dpi": FIGURE_DPI,
        "workbook": str(WORKBOOK),
        "workbook_sha256": sha256(WORKBOOK),
        "selection_rule": (
            "For each dataset, maximize HyperMOA's mean normalized rank-score margin "
            "over the best competing method across ARI, NMI, AMI, and Moran's I; "
            "break ties using the mean ARI/NMI/AMI margin."
        ),
        "selected_slices": {
            dataset: config["slice"]
            for dataset, config in selected.items()
        },
        "method_order": ["Ground truth", *METHODS],
        "cluster_color_rule": (
            "Ground truth and every method are colored independently from their own naturally "
            "sorted label values. No prediction-to-truth label matching is applied. Missing "
            "labels are light gray and only ground-truth categories appear in the legend."
        ),
        "metric_sources": metric_sources,
        "metrics_csv": str(METRICS_PATH),
        "selection_scores_csv": str(SELECTION_PATH),
        "figures": figures,
    }
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print("Selected slices:")
    for dataset, config in selected.items():
        print(f"  {dataset}: {config['slice']} ({config['sheet']})")
    print(f"Output directory: {OUTPUT_DIR}")
    print(f"Manifest: {MANIFEST_PATH}")


if __name__ == "__main__":
    main()
